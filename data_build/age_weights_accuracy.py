"""15-64歳の係数の作り方（Q3）の精度確認：B と B' の差を、東京都の勢圏で計算する。

- B ：勢圏（円）の 5歳刻み人口で、15〜29・30〜49・50〜64歳の重みを決める（いちばん細かい方法）。
- B'：勢圏の人口ではなく、勢圏がある市区町村全体の人口で重みを決める（今回採用する方法）。

どちらも、重みを掛けるのは「外出率 × トリップ数 × バス分担率」を区分ごとに掛けた値（全国都市交通特性調査 令和3年）。
差は次の2つで比べる：
  1. 15-64歳の「1人1日あたり」の値
  2. 勢圏全体の延べ乗車人員（4年齢階層の合計。15-64歳以外は B・B' で同じ）

使い方（プロジェクトのフォルダで）:
    .venv\\Scripts\\python data_build\\age_weights_accuracy.py
"""

from __future__ import annotations

import math

import geopandas as gpd
import numpy as np
import openpyxl
import pyogrio
from shapely.geometry import Point

from common import DATA_RAW, setup_console, vsizip
from sources import POPULATION

PREF = "13"
YEARS = (2025, 2040)
RADII = (300, 500)
PLANE_CRS = 6677  # JGD2011 平面直角座標系 第IX系（東京都）

# 確認に使う勢圏の中心（駅の位置など。区部と多摩）。city_type は調査の都市類型コード
# （三大都市圏の中心都市＝1、周辺都市１＝2。確認用の割り当てで、類型の判定ルールではない）
POINTS = [
    ("東京駅（千代田区）", 35.6812, 139.7671, 1),
    ("渋谷駅（渋谷区）", 35.6580, 139.7016, 1),
    ("池袋駅（豊島区）", 35.7295, 139.7109, 1),
    ("赤羽駅（北区）", 35.7780, 139.7210, 1),
    ("錦糸町駅（墨田区）", 35.6967, 139.8143, 1),
    ("蒲田駅（大田区）", 35.5625, 139.7161, 1),
    ("荻窪駅（杉並区）", 35.7043, 139.6200, 1),
    ("北千住駅（足立区）", 35.7497, 139.8049, 1),
    ("葛西駅（江戸川区）", 35.6644, 139.8715, 1),
    ("吉祥寺駅（武蔵野市）", 35.7030, 139.5797, 2),
    ("府中駅（府中市）", 35.6724, 139.4778, 2),
    ("立川駅（立川市）", 35.6980, 139.4137, 2),
    ("町田駅（町田市）", 35.5439, 139.4457, 2),
    ("多摩センター駅（多摩市）", 35.6246, 139.4240, 2),
    ("八王子駅（八王子市）", 35.6556, 139.3388, 2),
    ("青梅駅（青梅市）", 35.7880, 139.2756, 2),
]

SUB = ["15～29歳", "30～49歳", "50～64歳"]
SUB_PT = {"15～29歳": (4, 5, 6), "30～49歳": (7, 8, 9, 10), "50～64歳": (11, 12, 13)}
OTHER = {"0-14": ("5～14歳", (1, 2, 3)), "65-74": ("65～74歳", (14, 15)), "75+": ("75歳～", tuple(range(16, 21)))}

RAW = DATA_RAW / "mlit_r3"


def read_table(path, sheet, filt, cols):
    ws = openpyxl.load_workbook(path, read_only=True, data_only=True)[sheet]
    rows = list(ws.iter_rows(min_row=1, max_row=6, values_only=True))
    groups, names = rows[3], rows[5]
    idx = {(str(groups[i] or ""), str(n)): i for i, n in enumerate(names) if n is not None}
    plain = {n: i for (g, n), i in idx.items() if g == "" or n in ("平日／休日", "地域", "性別", "年齢階層", "世帯構成", "地域コード")}
    out = {}
    for r in ws.iter_rows(min_row=7, values_only=True):
        if r[0] is None or any(str(r[plain[k]]) != v for k, v in filt.items()):
            continue
        key = (int(r[plain["地域コード"]]), "weekday" if r[plain["平日／休日"]] == "平日" else "holiday", str(r[plain["年齢階層"]]))
        out[key] = {k: r[idx[c]] for k, c in cols.items()}
    return out


def coefficient_products() -> dict[tuple[int, str, str], float]:
    """（都市類型, 平日/休日, 調査の年齢階層）→ 外出率 × ネット原単位 × バス分担率（build_coefficients.py と同じ使い方）"""
    o = read_table(RAW / "001732066.xlsx", "4", {"性別": "合計", "世帯構成": "合計"},
                   {"out": ("外出率", "外出率"), "trip": ("ネット原単位", "ネット原単位")})
    m = read_table(RAW / "001732069.xlsx", "8-2", {"性別": "合計"}, {"bus": ("構成比（ネット）", "バス")})
    return {k: o[k]["out"] * o[k]["trip"] * m[k]["bus"] for k in o if k in m}


def main() -> None:
    setup_console()
    prod = coefficient_products()
    src = vsizip(DATA_RAW / "ksj" / POPULATION.file_name(PREF), POPULATION.inner(PREF))
    cols = ["MESH_ID", "SHICODE"] + [f"PT{c:02d}_{y}" for y in YEARS for c in range(1, 21)]
    df = pyogrio.read_dataframe(src, columns=cols)
    gdf = gpd.GeoDataFrame(df, geometry="geometry", crs=df.crs).to_crs(PLANE_CRS)
    area = gdf.geometry.area.values
    sindex = gdf.sindex
    print(f"メッシュ {len(gdf):,} 件。年次 {YEARS}、勢圏の半径 {RADII} m\n")


    results = []
    for name, lat, lng, city_type in POINTS:
        pt = gpd.GeoSeries([Point(lng, lat)], crs=4326).to_crs(PLANE_CRS).iloc[0]
        own = gdf[gdf.geometry.contains(pt)]
        if own.empty:
            own = gdf.iloc[[int(np.argmin(gdf.geometry.distance(pt).values))]]
        shi = own["SHICODE"].iloc[0]
        muni = gdf[gdf["SHICODE"] == shi]
        for radius in RADII:
            circle = pt.buffer(radius)
            cand = list(sindex.query(circle, predicate="intersects"))
            frac = np.array([gdf.geometry.iloc[i].intersection(circle).area / area[i] for i in cand])
            sub = gdf.iloc[cand]
            for year in YEARS:
                def pops(frame, w=None):
                    arr = {c: frame[f"PT{c:02d}_{year}"].astype(float).fillna(0).values for c in range(1, 21)}
                    mult = np.ones(len(frame)) if w is None else w
                    return {c: float((arr[c] * mult).sum()) for c in range(1, 21)}
                p_catch = pops(sub, frac)
                p_muni = pops(muni)
                for day in ("weekday", "holiday"):
                    v = {s: prod[(city_type, day, s)] for s in SUB}
                    def avg(p):
                        g = {s: sum(p[c] for c in SUB_PT[s]) for s in SUB}
                        tot = sum(g.values())
                        return sum(g[s] / tot * v[s] for s in SUB), tot, {s: g[s] / tot for s in SUB}
                    vB, popB, wB = avg(p_catch)
                    vBp, _, wBp = avg(p_muni)
                    # 勢圏全体の延べ乗車人員（15-64歳以外は B・B' で同じ）
                    others = sum(sum(p_catch[c] for c in pts) * prod[(city_type, day, key)] for _, (key, pts) in OTHER.items())
                    RB = others + popB * vB
                    RBp = others + popB * vBp
                    results.append((name, radius, year, day, vB, vBp, vBp / vB - 1, RBp / RB - 1, popB, wB, wBp, RB))
    # 出力
    print("勢圏ごとの結果（B に対する B' の差）")
    print(f"{'勢圏':<22}{'半径':>5}{'年':>6}{'日':>9}{'15-64の値 B':>13}{'B′':>9}{'差':>8}{'延べ乗車人員の差':>14}{'15-64人口':>10}")
    for r in results:
        print(f"{r[0]:<22}{r[1]:>5}{r[2]:>6}{r[3]:>9}{r[4]:>13.5f}{r[5]:>9.5f}{r[6] * 100:>+7.2f}%{r[7] * 100:>+13.2f}%{r[8]:>10.0f}")
    dv = np.array([r[6] for r in results])
    dr = np.array([r[7] for r in results])
    print("\n要約（絶対値の差）")
    for label, arr in (("15-64歳の1人1日あたり", dv), ("勢圏全体の延べ乗車人員", dr)):
        a = np.abs(arr) * 100
        print(f"  {label}: 平均 {a.mean():.2f}% / 中央値 {np.median(a):.2f}% / 最大 {a.max():.2f}% / 1%超 {int((a > 1).sum())}件 / 全 {len(a)}件")
    # 区部・多摩で分ける
    for label, pred in (("区部", lambda n: "区）" in n), ("多摩", lambda n: "市）" in n)):
        sel = [r for r in results if pred(r[0])]
        a1 = np.abs([r[6] for r in sel]) * 100
        a2 = np.abs([r[7] for r in sel]) * 100
        print(f"  [{label}] 15-64歳の値: 平均 {a1.mean():.2f}% 最大 {a1.max():.2f}% ／ 延べ乗車人員: 平均 {a2.mean():.2f}% 最大 {a2.max():.2f}%")
    # 平日・休日、年次ごと
    for label, pred in (("平日", lambda r: r[3] == "weekday"), ("休日", lambda r: r[3] == "holiday"),
                        (f"{YEARS[0]}年", lambda r: r[2] == YEARS[0]), (f"{YEARS[1]}年", lambda r: r[2] == YEARS[1])):
        sel = [r for r in results if pred(r)]
        a1 = np.abs([r[6] for r in sel]) * 100
        a2 = np.abs([r[7] for r in sel]) * 100
        print(f"  [{label}] 15-64歳の値: 平均 {a1.mean():.2f}% 最大 {a1.max():.2f}% ／ 延べ乗車人員: 平均 {a2.mean():.2f}% 最大 {a2.max():.2f}%")
    # 最大の差の勢圏
    worst = max(results, key=lambda r: abs(r[7]))
    print(f"\n延べ乗車人員の差が最大：{worst[0]} 半径{worst[1]}m {worst[2]}年 {worst[3]}  {worst[7] * 100:+.2f}%")
    print("  B の重み（勢圏）  ", {k: round(x, 3) for k, x in worst[9].items()})
    print("  B′ の重み（市区町村）", {k: round(x, 3) for k, x in worst[10].items()})


if __name__ == "__main__":
    main()
